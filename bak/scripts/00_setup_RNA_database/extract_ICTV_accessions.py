"""
Extract GenBank accessions for Riboviria from a VMR Excel file.
Handles single, labeled-segment (RNA1: ACC; RNA2: ACC), and unlabeled-list formats.
"""
import re
import sys
import openpyxl

sys.stdout = open(snakemake.log.log, "w")
sys.stderr = open(snakemake.log.err, "w")

vmr_xlsx      = snakemake.input[0]
out_txt       = snakemake.output[0]
sheet_name    = snakemake.params.sheet
realm_filter  = snakemake.params.realm_filter
cov_filter    = set(snakemake.params.coverage_filter)

# GenBank accession pattern: 1-6 uppercase letters, optional underscore, 5-9 digits, optional version
ACC_RE = re.compile(r'\b[A-Z]{1,6}_?[0-9]{5,9}(?:\.[0-9]+)?\b')


def parse_accessions(raw: str) -> list[str]:
    return ACC_RE.findall(raw)


wb = openpyxl.load_workbook(vmr_xlsx, read_only=True, data_only=True)
ws = wb[sheet_name]

headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
realm_idx = headers.index("Realm")
cov_idx   = headers.index("Genome coverage")
acc_idx   = headers.index("Virus GENBANK accession")

accessions = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[realm_idx] == realm_filter and row[cov_idx] in cov_filter:
        raw = row[acc_idx]
        if raw:
            accessions.extend(parse_accessions(str(raw)))

unique_accessions = sorted(set(accessions))
print(f"Extracted {len(unique_accessions)} unique accessions", file=sys.stderr)

with open(out_txt, "w") as f:
    f.write("\n".join(unique_accessions) + "\n")
